#!/usr/bin/env python3
"""TouchRetouch ASA — experiment briefs + measurement/analytics plan (doc)."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

INK=RGBColor(0x1F,0x24,0x30); ACC=RGBColor(0x2E,0x5A,0xAC); ACC2=RGBColor(0x0E,0x7C,0x61); GREY=RGBColor(0x6B,0x72,0x80)
d=Document()
st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5); st.font.color.rgb=INK

def h1(t):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=ACC; p.space_after=Pt(4); return p
def h2(t):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=INK
    p.space_before=Pt(10); p.space_after=Pt(3); return p
def para(t,italic=False,color=INK,size=10.5):
    p=d.add_paragraph(); r=p.add_run(t); r.italic=italic; r.font.color.rgb=color; r.font.size=Pt(size); return p
def bullet(t,bold_lead=None):
    p=d.add_paragraph(style="List Bullet")
    if bold_lead:
        r=p.add_run(bold_lead+": "); r.bold=True
    p.add_run(t); return p

h1("TouchRetouch — ASA Experiment Briefs & Measurement Plan")
para("Companion to the strategy deck and the ROAS model workbook. Covers (A) the prioritised experiment backlog with full briefs, and (B) the measurement / attribution plan for iOS under ATT.", italic=True, color=GREY)

# ---- assumptions box ----
h2("Model anchor (read first)")
bullet("Revenue is modelled on first-year ARPPU ($22.2, net of Apple fee). Rev/install = 15% × 20% × $22.2 = $0.666.", "Basis")
bullet("This is the SAME basis as the task's '~100% ROAS at month 1' anchor, so 'ROAS (gate)' in the model is comparable to the 125% requirement.", "Why")
bullet("Real cash recognised by day 30 is a fraction of Y1 — the 'D30 cash ROAS' column is informational until we replace the 0.45 placeholder with RevenueCat's real D30/Y1 curve.", "Caveat")

# =================== PART A ===================
h1("Part A — Experiment backlog (prioritised)")
para("Prioritised by revenue unlock × confidence × ease (ICE). Each brief states hypothesis, design, target KPI, minimum runtime, and success/kill criteria. Kill fast, scale winners.", italic=True, color=GREY)

exps=[
 ("EXP-1 · Brand bid cap + incrementality pause test  (Priority: HIGHEST)",
  "Brand is 58% of spend with max bids €25–35 against actual CPTs of €1.2–1.9. Much of this is organic traffic we already win #1 for. Capping bids and running a pause test reveals how much brand paid spend is truly incremental.",
  "Cap brand max CPT to ~€2–3. Then pause the brand campaign for 1–2 weeks and measure how much of the lost paid brand installs are recovered by organic (branded search). Use a clean before/after window; watch competitor brand-bidding.",
  "Brand incremental installs; total (paid+organic) brand installs; blended CPI.",
  "2 weeks pause (min) + 1 week read.",
  "SUCCESS: organic backfills ≥70% of paused brand installs → keep brand bids capped permanently, reallocate freed budget to generic. KILL: organic recovers <50% OR a competitor floods your brand term → restore a capped defensive bid."),

 ("EXP-2 · Search Match isolation + negative flow  (Priority: HIGH)",
  "Search Match (31% of spend) runs inside generic campaigns with exact keywords and no negatives, so it cannibalises exact terms and runs semi-blind on Maximize Conversions.",
  "Move Search Match into its own Discovery campaign. Add all Core exact keywords as exact negatives to it. Harvest converting search terms into Core exact ad groups weekly.",
  "Generic CPA; wasted spend %; net new keywords harvested.",
  "3–4 weeks (SKAN postback lag).",
  "SUCCESS: generic CPA −15% at equal or higher install volume → make permanent. KILL: install volume drops >20% with no CPA gain → revert / rebalance budgets."),

 ("EXP-3 · Competitor conquesting in Search Results  (Priority: HIGH)",
  "The competitor campaign runs only on Product Pages (CPA €17 vs €0.75–2 elsewhere) and there is no keyword-level conquesting in Search Results — the highest-intent competitor placement.",
  "Launch a Competitor campaign on Search Results with exact competitor terms (facetune, snapseed, picsart, lightroom, youcam). Pair with the 'Competitor Switch' CPP. Rework or cut the Product Pages competitor.",
  "CPA(trial) from competitor; incremental installs; Trial→Paid on competitor cohort.",
  "3–4 weeks.",
  "SUCCESS: CPA(trial) ≤ target with positive incremental installs → scale. KILL: CPA > generic × 1.5 with no volume benefit → pause, keep only best-performing terms."),

 ("EXP-4 · CPP 'Object Cleanup' vs default page  (Priority: MEDIUM-HIGH)",
  "Generic/feature 'remove object' intent is served by the default page, not a tailored one. A CPP matched to the search intent should lift tap→install CR.",
  "Serve the 'Object Cleanup' CPP (before/after object-removal screenshots, one-tap erase preview) to the object-removal / magic-eraser ad groups via ASA ad variations. Compare against the default page on the same terms.",
  "tap→install CR; trial-start rate; CPA(trial) by page.",
  "Until ~9,000 impressions/variant or 2 weeks (for ~10% lift at 80% power).",
  "SUCCESS: +10% CR (stat-sig) → ship CPP to those ad groups; build next CPP. KILL: no lift after target impressions → iterate creative or retire."),

 ("EXP-5 · PPO screenshot A/B on the default listing  (Priority: MEDIUM)",
  "The default store page drives BOTH organic and paid-to-default conversion; a screenshot/icon test lifts install rate across all traffic.",
  "Run a Product Page Optimization test: up to 3 treatments vs default (lead screenshot = finished before/after; alt hero; icon variant). Apple splits traffic and reports statistically.",
  "Store install rate (impression→install); downstream trial rate.",
  "Until statistical significance (Apple-managed), typically 2–4 weeks.",
  "SUCCESS: +5% install rate (stat-sig) → promote winner to default. KILL: no lift → keep default, test a bolder creative direction."),

 ("EXP-6 · Geo split of T1_Generic (isolate US)  (Priority: MEDIUM)",
  "T1_Generic mixes ~40 storefronts (US + UZ/KG) under one budget and bid; cheap junk geos (TTR 4–5%) eat budget that should be US.",
  "Split T1_Generic by geo tier; isolate US into its own campaign with its own bids; move low-value geos to separate small-budget campaigns with their own targets, or pause.",
  "US CPI; blended CPI; wasted spend on sub-target geos.",
  "2–3 weeks.",
  "SUCCESS: US-isolated CPI ≤ target and blended CPI improves → make permanent. KILL: n/a (structural hygiene; keep unless US volume collapses)."),
]
for title,hyp,design,kpi,runtime,crit in exps:
    h2(title)
    bullet(hyp,"Hypothesis")
    bullet(design,"Design")
    bullet(kpi,"Target KPI")
    bullet(runtime,"Min runtime")
    bullet(crit,"Success / Kill")

# =================== PART B ===================
d.add_page_break()
h1("Part B — Measurement & attribution plan")

h2("B.1 — The 30% console-vs-MMP install gap")
para("The ASA console currently shows ~30% more installs than AppsFlyer. Likely causes:")
bullet("AdServices (console) counts redownloads and view-through conversions, and uses Apple's own attribution window.","Over-counting side")
bullet("AppsFlyer dedupes, applies its own attribution window, and only counts installs where its SDK fires on first open — so ATT-declined / SKAN-null / delayed-postback installs and uninstalled-before-open users drop out.","Under-counting side")
bullet("Tap-through vs view-through counting differences, timezone/window mismatches, and SKAN postback delay/aggregation add noise.","Structural")

h2("B.2 — Reconciliation approach")
bullet("Agree ONE attribution window across console and MMP.")
bullet("Use the AdServices API as the source of truth for ASA installs — it is deterministic and survives ATT because Apple owns both the store and the ad network.")
bullet("Build a documented weekly bridge: console installs − redownloads − ATT/SKAN-null gaps ≈ MMP attributed installs. Track the residual %; a sudden change flags a tracking break, not a performance change.")

h2("B.3 — Source of truth by metric")
tbl=d.add_table(rows=1,cols=3); tbl.style="Light Grid Accent 1"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
hdr=tbl.rows[0].cells
for i,t in enumerate(["System","Owns these metrics","Notes"]):
    hdr[i].paragraphs[0].add_run(t).bold=True
rows=[
 ("ASA console","Spend, impressions, TTR, CPT, taps, AdServices installs","First-party; deterministic ASA install truth under ATT"),
 ("AppsFlyer (MMP)","Attributed installs + downstream events across channels","Cross-channel view; blends SKAN + AdServices"),
 ("RevenueCat / Amplitude","Trials, payers, revenue, ARPPU, LTV, D30/Y1 curve","Revenue truth; feeds the ROAS model"),
]
for a,b,c in rows:
    cells=tbl.add_row().cells
    cells[0].paragraphs[0].add_run(a); cells[1].paragraphs[0].add_run(b); cells[2].paragraphs[0].add_run(c)

h2("B.4 — Reading iOS data under ATT / AdAttributionKit")
bullet("ASA installs: AdServices deterministic attribution — reliable at keyword level even when ATT is declined.")
bullet("Downstream (trial/purchase): SKAN / AdAttributionKit conversion values. Design a CV schema encoding funnel milestones (install → trial → purchase) or revenue buckets; RevenueCat fires the value on subscription events.")
bullet("Constraints to respect: postback delay + conversion windows (0–2 / 3–7 / 8–35 days), aggregation (no user-level), and crowd-anonymity thresholds that null-out low-volume sources. Keyword-level revenue is therefore approximated by narrow ad-group / campaign structure so campaign-level SKAN ≈ keyword-level.")
bullet("Where SKAN is null (small keywords): use MMP predictive/blended modelling and evaluate by cohort (D7 leading, D30 gate), not by same-day installs.")

h2("B.5 — Weekly report KPIs")
for k in ["Spend","Impressions & TTR","CR (tap→install)","CPT & CPI","Installs — console vs MMP delta %",
          "CPA(trial) and CPA(payer)","Trial→Paid %","ROAS — D7 (leading) and D30 (gate)",
          "Organic keyword rank on top terms (paid–organic sync)","Organic install share","Share of defended brand traffic"]:
    bullet(k)

out="/home/user/asotraff/outputs/touchretouch-testtask/TouchRetouch_Experiments_and_Measurement.docx"
d.save(out); print("saved",out)
