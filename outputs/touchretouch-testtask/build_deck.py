#!/usr/bin/env python3
"""TouchRetouch ASA test-task deck (<=12 slides)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK=RGBColor(0x1F,0x24,0x30); ACC=RGBColor(0x2E,0x5A,0xAC); ACC2=RGBColor(0x0E,0x7C,0x61)
WARN=RGBColor(0xB5,0x76,0x1A); GREY=RGBColor(0x6B,0x72,0x80); WHITE=RGBColor(0xFF,0xFF,0xFF)
LIGHT=RGBColor(0xEA,0xF0,0xFB); BG=RGBColor(0xF7,0xF8,0xFA)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=prs.slide_width,prs.slide_height

def slide():
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(1,0,0,SW,SH); r.fill.solid(); r.fill.fore_color.rgb=BG; r.line.fill.background()
    r.shadow.inherit=False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2,r._element)
    return s

def box(s,l,t,w,h):
    tb=s.shapes.add_textbox(l,t,w,h); tb.text_frame.word_wrap=True; return tb.text_frame

def setpar(p,text,size,color=INK,bold=False,align=PP_ALIGN.LEFT,space=6,italic=False,font="Calibri"):
    p.text=text; p.alignment=align; p.space_after=Pt(space)
    r=p.runs[0]; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color; r.font.name=font

def title_bar(s,eyebrow,title):
    bar=s.shapes.add_shape(1,0,0,SW,Inches(1.35)); bar.fill.solid(); bar.fill.fore_color.rgb=ACC
    bar.line.fill.background(); bar.shadow.inherit=False
    tf=box(s,Inches(0.55),Inches(0.18),Inches(12.3),Inches(1.0))
    setpar(tf.paragraphs[0],eyebrow.upper(),11,WHITE,bold=True,space=2)
    p=tf.add_paragraph(); setpar(p,title,25,WHITE,bold=True,space=0)

def bullets(s,items,l,t,w,h,size=15,gap=8):
    tf=box(s,l,t,w,h)
    for i,(txt,lvl,*rest) in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        col=rest[0] if rest else (INK if lvl==0 else GREY)
        bold=(lvl==0)
        prefix="●  " if lvl==0 else "–  "
        setpar(p,prefix+txt,size if lvl==0 else size-2,col,bold=bold,space=gap)
        p.level=0 if lvl==0 else 1

def chip(s,l,t,w,text,color=ACC2,h=Inches(0.42)):
    c=s.shapes.add_shape(5,l,t,w,h); c.fill.solid(); c.fill.fore_color.rgb=color; c.line.fill.background(); c.shadow.inherit=False
    tf=c.text_frame; tf.word_wrap=True; tf.margin_top=Pt(2); tf.margin_bottom=Pt(2)
    setpar(tf.paragraphs[0],text,11,WHITE,bold=True,align=PP_ALIGN.CENTER,space=0)
    return c

def table(s,data,l,t,w,colw,rowh=Inches(0.4),head=ACC,fs=11):
    rows=len(data); cols=len(data[0])
    gt=s.shapes.add_table(rows,cols,l,t,w,rowh*rows).table
    for ci,cw in enumerate(colw): gt.columns[ci].width=cw
    for ri,row in enumerate(data):
        for cidx,val in enumerate(row):
            cell=gt.cell(ri,cidx); cell.margin_left=Pt(5); cell.margin_right=Pt(4)
            cell.margin_top=Pt(2); cell.margin_bottom=Pt(2)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            para=cell.text_frame.paragraphs[0]
            rr=para.add_run(); rr.text=str(val); rr.font.size=Pt(fs); rr.font.name="Calibri"
            if ri==0:
                cell.fill.solid(); cell.fill.fore_color.rgb=head
                rr.font.bold=True; rr.font.color.rgb=WHITE
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb=WHITE if ri%2 else RGBColor(0xF0,0xF3,0xF7)
                rr.font.color.rgb=INK
    return gt

# ---------- SLIDE 1: TITLE ----------
s=slide()
bar=s.shapes.add_shape(1,0,Inches(2.4),SW,Inches(2.7)); bar.fill.solid(); bar.fill.fore_color.rgb=ACC; bar.line.fill.background(); bar.shadow.inherit=False
tf=box(s,Inches(0.7),Inches(2.65),Inches(12),Inches(2.3))
setpar(tf.paragraphs[0],"APPLE SEARCH ADS · TEST TASK",13,WHITE,bold=True,space=6)
p=tf.add_paragraph(); setpar(p,"TouchRetouch — bring ASA in-house, fix it, scale to $70K/mo @ ≥125% ROAS",26,WHITE,bold=True,space=6)
p=tf.add_paragraph(); setpar(p,"Photo editor (iOS) · freemium weekly trial → subscription · US market · 2-quarter plan",14,WHITE,space=0)
tf2=box(s,Inches(0.7),Inches(5.4),Inches(12),Inches(1))
setpar(tf2.paragraphs[0],"Audit → restructure → measurement fix → CPP/ASO → scale. Base plan clears 125% on CPT efficiency alone; funnel gains are upside.",13,GREY,italic=True,space=0)

# ---------- SLIDE 2: AUDIT ----------
s=slide(); title_bar(s,"1 · Account audit","Key problems, ranked by revenue impact")
data=[["#","Problem","Why it's money","Fix"],
["1","Brand = 58% of spend (~$20K) + cannibalisation","Paying for installs organic would win free; max bid €25–35 vs €1.2–1.9 actual","Cap bids ~€2–3; brand-pause incrementality test; reallocate"],
["2","Search Match (31%) inside generic, no negatives","SM competes with own exact keywords, runs semi-blind","Isolate SM in Discovery; Core keywords as negatives; harvest"],
["3","T1_Generic: 40 storefronts, 1 budget/bid","US competes with UZ/KG junk geos (TTR 4–5%)","Split by geo/tier; isolate US; pause junk"],
["4","Competitor only on Product Pages (CPA €17)","10× inefficient; missing Search Results conquesting","Launch keyword competitor in Search Results"],
["5","Feature campaign: single-digit installs @ €4–7","Sub-scale, overbid on low volume","Fold into generic ad group"],
["6","Console vs MMP installs differ 30%","Can't trust ROAS decisions — foundation","Reconcile (see measurement slide)"]]
table(s,data,Inches(0.4),Inches(1.6),Inches(12.5),[Inches(0.4),Inches(3.6),Inches(4.6),Inches(3.9)],rowh=Inches(0.78),fs=11)

# ---------- SLIDE 3: ARCHITECTURE ----------
s=slide(); title_bar(s,"1 · Restructure","Target account architecture")
cols=[("Brand (US)","Exact · bid capped ~€2–3","Defend cheaply, stop overpay; measure incrementality",ACC),
("Generic Core (US)","Exact · manual bids","Proven head/mid terms under CPA target",ACC2),
("Generic Discovery","Search Match + Broad, ISOLATED","Find new terms → harvest to Core; Core as negatives",ACC2),
("Competitor (US)","Exact · Search Results","Conquesting w/ own CPP (was PP-only, CPA €17)",WARN),
("Feature","Ad group inside Generic","No separate zoo of tiny-volume terms",GREY)]
y=Inches(1.7)
for name,cfg,why,c in cols:
    chip(s,Inches(0.5),y,Inches(3.0),name,c,h=Inches(0.55))
    tf=box(s,Inches(3.7),y-Inches(0.02),Inches(9.2),Inches(0.7))
    setpar(tf.paragraphs[0],cfg,13,INK,bold=True,space=2)
    p=tf.add_paragraph(); setpar(p,why,12,GREY,space=0)
    y+=Inches(0.92)
tf=box(s,Inches(0.5),Inches(6.5),Inches(12.3),Inches(0.8))
setpar(tf.paragraphs[0],"Negatives flow: Core keywords → negatives in Discovery/Search Match · Brand terms → negatives in Generic & Competitor. Placements: Search Results primary; Search/Today tab tested; Product Pages only where it earns.",12,INK,italic=True,space=0)

# ---------- SLIDE 4: KEYWORDS ----------
s=slide(); title_bar(s,"2 · Keyword strategy","Research + sample pool by intent")
tf=box(s,Inches(0.5),Inches(1.5),Inches(12.3),Inches(0.6))
setpar(tf.paragraphs[0],"Process: product seeds + App Store autosuggest + AppTweak/ASOMobile (volume/difficulty/intent) + ASA search-term reports + competitor metadata + Discovery harvest. Re-research quarterly.",12,INK,space=0)
data=[["Brand","Generic","Competitor","Feature"],
["touchretouch (exact)","photo editor (exact)","facetune (exact)","magic eraser (exact)"],
["touchretouch app (exact)","object remover (exact)","snapseed (exact)","remove people (exact)"],
["","remove object from photo","picsart (exact)","background eraser (exact)"],
["","photo retouch (exact)","lightroom (exact)","unwanted object remover"],
["","unblur / photo enhancer","youcam perfect (exact)","object eraser (exact)"]]
table(s,data,Inches(0.5),Inches(2.3),Inches(12.3),[Inches(3.0),Inches(3.3),Inches(3.0),Inches(3.0)],rowh=Inches(0.5),fs=11)
tf=box(s,Inches(0.5),Inches(5.9),Inches(12.3),Inches(1.2))
setpar(tf.paragraphs[0],"Match logic: exact for known performers; Broad+Search Match isolated in Discovery to harvest. Competitor brands: BID in ASA, but NEVER in metadata (trademark / 2.3.7).",12,WARN,bold=True,space=4)

# ---------- SLIDE 5: ASO / paid-organic ----------
s=slide(); title_bar(s,"2 · ASO & paid–organic sync","Protect and grow organic rank")
bullets(s,[
("Metadata: top generic head → Subtitle; feature terms → Keyword field; Title = TouchRetouch + top category term.",0),
("Competitor brands: bid in ASA only — keep them OUT of metadata (trademark).",1),
("Halo: paid install velocity lifts organic rank on the same term → free installs on top.",0),
("Cannibalisation: bidding own brand partly pays for installs organic would win free.",1),
("Measure paid→organic: track organic rank per keyword (AppTweak) before/after paid push.",0),
("Incrementality: brand-pause test + geo-holdout to isolate true incremental paid.",1),
("Default listing IS an ASO surface: PPO A/B on screenshots/icon to lift install rate; CPP complements it per intent.",0),
],Inches(0.5),Inches(1.6),Inches(12.3),Inches(5),size=15,gap=11)

# ---------- SLIDE 6: CPP ----------
s=slide(); title_bar(s,"2 · Custom Product Pages","3 concepts paired to keyword themes")
data=[["CPP","Serves keywords","Page changes","Uplift metric"],
["Object Cleanup","remove object, magic eraser, cleanup","Before/after object-removal screens; one-tap erase preview","tap→install CR + trial-start rate"],
["AI Enhance / Retouch","photo enhancer, ai editor, retouch","Portrait retouch/enhance leads; 'AI in one tap' msg","CR + trial rate"],
["Competitor Switch","facetune, snapseed, picsart","'Everything they do, faster'; familiar UI cues","CR from competitor camps + trial→paid"]]
table(s,data,Inches(0.5),Inches(1.7),Inches(12.3),[Inches(2.4),Inches(3.2),Inches(4.2),Inches(2.5)],rowh=Inches(0.95),fs=12)
tf=box(s,Inches(0.5),Inches(5.4),Inches(12.3),Inches(1))
setpar(tf.paragraphs[0],"Judged primarily on tap→install CR (what a CPP influences) and downstream trial-start; ultimately CPA(trial)/ROAS by CPP. Each CPP paired to its ad group so creative matches search intent.",12,GREY,italic=True,space=0)

# ---------- SLIDE 7: ROAS MODEL ----------
s=slide(); title_bar(s,"3 · Budget & ROAS model","Path from $35K → $70K @ ≥125%")
data=[["Phase","Spend","CPI","Installs","Revenue","ROAS","Gate"],
["M0 Baseline","$35K","$0.666","52.6K","$35.0K","100%","—"],
["M1 Bank efficiency","$35K","$0.560","62.5K","$41.6K","119%","build"],
["M2 Gate cleared","$40K","$0.531","75.3K","$50.2K","125%","PASS"],
["M3 Scale","$55K","$0.528","104K","$69.4K","126%","PASS"],
["Q2 Target","$70K","$0.525","133K","$88.8K","127%","PASS"]]
table(s,data,Inches(0.5),Inches(1.7),Inches(9.0),[Inches(2.2),Inches(1.1),Inches(1.2),Inches(1.3),Inches(1.4),Inches(1.0),Inches(0.9)],rowh=Inches(0.55),fs=12)
tf=box(s,Inches(9.8),Inches(1.7),Inches(3.2),Inches(4.5))
setpar(tf.paragraphs[0],"Unit economics",13,ACC,bold=True,space=4)
for t in ["Rev/install (Y1) = 3% × $22.2 = $0.666","Break-even CPI = $0.666","Target CPI (125%) = $0.533"]:
    p=tf.add_paragraph(); setpar(p,t,12,INK,space=4)
p=tf.add_paragraph(); setpar(p,"Bank efficiency FIRST (M1), then scale into the headroom. Base plan hits 125% on CPI alone — funnel gains are upside.",12,GREY,italic=True,space=4)
tf2=box(s,Inches(0.5),Inches(5.6),Inches(12.3),Inches(1.3))
setpar(tf2.paragraphs[0],"Bidding: start from target CPA (CPT = CPA × CR); move on downstream (trial/payer, not installs); brand capped at actual CPT.",13,INK,bold=True,space=4)
p=tf2.add_paragraph(); setpar(p,"Sensitivity: ROAS chain is multiplicative → every metric has equal % elasticity. Prioritise by headroom × control: CPT efficiency (fastest), Trial→Paid (most headroom, 20%), tap→install CR (CPP/PPO).",12,WARN,bold=True,space=0)

# ---------- SLIDE 8: MEASUREMENT ----------
s=slide(); title_bar(s,"4 · Measurement & attribution","Reconcile the 30% gap; source of truth")
bullets(s,[
("30% console>MMP causes: AdServices counts redownloads & view-through, different window; MMP dedupes, drops ATT/SKAN-null installs, needs SDK first-open.",0),
("Reconcile: agree ONE attribution window; AdServices API = source of truth for ASA installs (deterministic, survives ATT); documented console→MMP bridge, weekly.",0),
("Source of truth — ASA console: spend, TTR, CPT, taps + AdServices installs.",1),
("MMP (AppsFlyer): attributed installs + downstream across channels.",1),
("RevenueCat / Amplitude: trials, payers, revenue, LTV, D30/Y1 curve.",1),
("Under ATT: AdServices deterministic for ASA installs; SKAN conversion values for downstream (non-ATT); predictive/blended for revenue.",0),
("Weekly KPIs: Spend, TTR, CR, CPT, CPI, installs (console vs MMP Δ%), CPA(trial), CPA(payer), Trial→Paid%, ROAS (D7 lead + D30 gate), organic rank of top terms, organic install share.",0),
],Inches(0.5),Inches(1.6),Inches(12.3),Inches(5.3),size=13.5,gap=10)

# ---------- SLIDE 9: EXPERIMENTS ----------
s=slide(); title_bar(s,"5 · Experiments backlog","Prioritised by $-unlock")
data=[["Test","Target KPI","Success / Kill"],
["1 · Brand bid cap + pause incrementality","Brand incremental installs","Organic backfills ≥70% → cut permanently / Kill: organic −>30%"],
["2 · Search Match isolation + negatives","Generic CPA","CPA −15% at equal volume / Kill: volume −>20%"],
["3 · Competitor in Search Results","CPA(trial)","CPA ≤ target, incremental+ / Kill: CPA > generic ×1.5"],
["4 · CPP 'Object Cleanup' vs default","tap→install CR + trial rate","+10% CR stat-sig / Kill: no lift after N impressions"],
["5 · PPO screenshot test (default)","Store install rate","+5% stat-sig / Kill: no lift"],
["6 · Geo-split T1 (isolate US)","US CPI","US CPI ≤ target / Kill: —"]]
table(s,data,Inches(0.5),Inches(1.7),Inches(12.3),[Inches(4.5),Inches(3.0),Inches(4.8)],rowh=Inches(0.7),fs=12)

# ---------- SLIDE 10: 90-DAY ----------
s=slide(); title_bar(s,"6 · 90-day channel plan","0 → 30 → 60 → 90")
cols=[("Days 0–30 · Stabilise",["Agency transition: access, history, search-term reports","Audit + fix measurement (AdServices/SKAN/MMP)","Quick wins: cap brand bids, isolate SM, pause junk geos","Ship 1–2 low-risk fixes; stand up reporting"],ACC),
("Days 30–60 · Rebuild",["Live architecture: Brand/Core/Discovery/Competitor-SR/Feature","First CPP + PPO tests","Paid–organic keyword sync","Expand generic + competitor"],ACC2),
("Days 60–90 · Scale",["Scale into banked efficiency → ~$55K @ ≥125%","Validate CPP winners","Q2 scaling plan → $70K (budget, keywords, CPPs, geos)"],WARN)]
x=Inches(0.4)
for name,items,c in cols:
    chip(s,x,Inches(1.6),Inches(4.1),name,c,h=Inches(0.55))
    tf=box(s,x,Inches(2.3),Inches(4.1),Inches(4.5))
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        setpar(p,"•  "+it,13,INK,space=10)
    x+=Inches(4.3)
tf=box(s,Inches(0.4),Inches(6.6),Inches(12.5),Inches(0.7))
setpar(tf.paragraphs[0],"Need from team: creative for 2–3 CPPs (screens + preview) · tools (AppTweak/ASOMobile, opt. SplitMetrics) · budget authority + ROAS sign-off · CMO decisions on brand-defense risk & geo expansion · product collab on paywall (Trial→Paid).",12,GREY,italic=True,space=0)

# ---------- SLIDE 11: RISKS ----------
s=slide(); title_bar(s,"7 · Risks & assumptions","What the model depends on")
bullets(s,[
("#1 assumption: revenue modelled on first-year ARPPU (same basis as the '~100%' anchor). Real cash by D30 = ROAS × D30/Y1 recognition — pull the curve from RevenueCat.",0),
("Brand incrementality: assumes organic backfills most paused brand installs. Risk: competitors bid your brand. Mitigate: capped defensive bid + monitor.",0),
("Scaling: assumes generic CPT ~stable as spend 2×. Risk: auction inflation / limited US volume → CPI rises → ROAS breaks. Mitigate: widen keywords, CPP/PPO CR gains, geo expansion for volume.",0),
("Measurement: SKAN/ATT noise & delay. Mitigate: AdServices + documented reconciliation.",0),
("External: seasonality, competitor bidding, ATT opt-in rate. Monitor and re-forecast monthly.",0),
],Inches(0.5),Inches(1.7),Inches(12.3),Inches(5),size=15,gap=16)

# ---------- SLIDE 12: CLOSE ----------
s=slide()
bar=s.shapes.add_shape(1,0,Inches(2.6),SW,Inches(2.3)); bar.fill.solid(); bar.fill.fore_color.rgb=ACC2; bar.line.fill.background(); bar.shadow.inherit=False
tf=box(s,Inches(0.7),Inches(2.85),Inches(12),Inches(1.9))
setpar(tf.paragraphs[0],"The thesis in one line",13,WHITE,bold=True,space=8)
p=tf.add_paragraph(); setpar(p,"Cut brand + geo waste to bank ROAS headroom first, then scale efficient US generic & competitor into it — hitting $70K @ ≥125% by end of Q2.",20,WHITE,bold=True,space=0)
tf2=box(s,Inches(0.7),Inches(5.3),Inches(12),Inches(1))
setpar(tf2.paragraphs[0],"Deliverables: this deck · ROAS model (xlsx, editable) · experiment briefs + measurement plan (doc).",13,GREY,italic=True,space=0)

out="/home/user/asotraff/outputs/touchretouch-testtask/TouchRetouch_ASA_Deck.pptx"
prs.save(out); print("saved",out,"slides:",len(prs.slides._sldIdLst))
